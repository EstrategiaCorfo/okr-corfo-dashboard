#!/usr/bin/env python3
"""Convert the strategy and reporting sheets to the v6 dashboard data files."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re

import openpyxl


def txt(value):
    return "" if value is None else str(value).strip()


def number(value, label):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"{label}: expected a number")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}: invalid number {value!r}") from exc


def iso_date(value, label):
    if value is None or value == "":
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(txt(value), fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"{label}: invalid date {value!r}")


def periods(value):
    raw = txt(value)
    matches = list(re.finditer(r"(?:(20\d{2})\s*/\s*)?Q([1-4])", raw))
    if not matches or not matches[0].group(1):
        raise ValueError(f"Invalid target period: {raw!r}")
    year = int(matches[0].group(1))
    ordinals = []
    for match in matches:
        year = int(match.group(1) or year)
        ordinals.append(year * 4 + int(match.group(2)) - 1)
    if len(ordinals) == 2 and "-" in raw:
        if ordinals[1] < ordinals[0]:
            raise ValueError(f"Reversed target period: {raw!r}")
        ordinals = list(range(ordinals[0], ordinals[1] + 1))
    return list(dict.fromkeys(f"{p // 4}/Q{p % 4 + 1}" for p in ordinals))


def header(sheet, row, col, expected):
    actual = txt(sheet.cell(row, col).value).split("\n", 1)[0].strip()
    if actual != expected:
        raise ValueError(f"{sheet.title}!{row}:{col}: expected {expected!r}, got {actual!r}")


def build(input_path, strategy_sheet, tracking_sheet, confirm_ge):
    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    try:
        strategy_ws, tracking_ws = workbook[strategy_sheet], workbook[tracking_sheet]
        for col, expected in [(2, "Objetivo Estratégico"), (6, "ID KR"), (7, "Nombre del KR"), (8, "Trimestre meta KR"), (12, "Producto del KR (Referencial)"), (14, "Owner (Lider y responsable)")]:
            header(strategy_ws, 4, col, expected)
        for col, expected in [(1, "ID KR"), (3, "Periodo de reporte"), (6, "Instancia"), (7, "Instancia del registro"), (8, "Estado de ejecución"), (12, "Métrica (%)"), (13, "Nivel de confianza (semáforo)"), (19, "Periodo de traspaso")]:
            header(tracking_ws, 2, col, expected)

        objectives, notes, catalogue = {}, {}, []
        ids = set()
        for row in strategy_ws.iter_rows(min_row=5):
            v = [c.value for c in row]
            identifier = txt(v[5])
            if not identifier:
                continue
            if not re.fullmatch(r"KR-\d{3,}", identifier) or identifier in ids:
                raise ValueError(f"Invalid or duplicate KR ID: {identifier}")
            ids.add(identifier)
            objective = txt(v[1])
            match = re.match(r"Objetivo\s+(\d+)\s*:", objective)
            if not match:
                raise ValueError(f"{identifier}: invalid objective {objective!r}")
            obj = objectives.setdefault(objective, {"name": objective, "number": int(match.group(1)), "dimension": txt(v[0]), "vision": txt(v[2]), "subobjectives": []})
            subobjective = txt(v[4])
            if subobjective not in obj["subobjectives"]:
                obj["subobjectives"].append(subobjective)
            justification = txt(v[3])
            if objective in notes and notes[objective] != justification:
                raise ValueError(f"{identifier}: inconsistent objective justification")
            notes[objective] = justification
            catalogue.append({"id": identifier, "name": txt(v[6]), "objective": objective,
                "subobjective": subobjective, "dimension": txt(v[0]), "vision": txt(v[2]),
                "target_period": txt(v[7]), "target_year": txt(v[8]), "quarter": txt(v[9]),
                "periods": periods(v[7]), "target_date": iso_date(v[10], identifier + " target date"),
                "product": txt(v[11]), "verification": "", "area": txt(v[12]),
                "owner": txt(v[13]), "substitutes": txt(v[14]), "comments": "",
                "source_sheet": strategy_sheet, "source_row": row[0].row})
        by_id = {k["id"]: k for k in catalogue}
        confidence = {"Alto (On Track)": "On Track", "Medio (Off Track)": "Off Track", "Bajo (At Risk)": "At Risk", "No iniciado": "No iniciado", "": ""}
        instances = {"Apertura de Q": "Apertura", "Apertura": "Apertura", "Revisión intermedia": "Revisión intermedia", "Cierre de Q": "Cierre", "Cierre": "Cierre", "Pendiente": "Pendiente"}
        records, seen = [], set()
        for row in tracking_ws.iter_rows(min_row=3):
            v = [c.value for c in row]
            identifier = txt(v[0])
            if not identifier:
                continue
            if identifier not in by_id:
                raise ValueError(f"Reporting KR is absent from catalogue: {identifier}")
            source_name = txt(v[1])
            if source_name and " ".join(source_name.split()) != " ".join(by_id[identifier]["name"].split()):
                raise ValueError(f"{identifier}: reporting name does not match catalogue")
            period, instance, source_moment, execution, status = txt(v[2]), txt(v[5]), txt(v[6]), txt(v[7]), txt(v[12])
            if not re.fullmatch(r"20\d{2}/Q[1-4]", period):
                raise ValueError(f"{identifier}: invalid reporting period {period!r}")
            if instance not in instances or execution not in {"No iniciado", "En proceso", "Completado"} or status not in confidence:
                raise ValueError(f"{identifier}: invalid instance, execution or confidence")
            if source_moment not in {"Propuesta Owner", "Confirmado GE"}:
                raise ValueError(f"{identifier}: invalid record validation {source_moment!r}")
            moment = "Confirmado GE" if confirm_ge else source_moment
            report_date = iso_date(v[3], identifier + " report date")
            key = (identifier, period, instances[instance], moment, report_date)
            if key in seen:
                raise ValueError(f"Duplicate reporting record: {key}")
            seen.add(key)
            progress = number(v[11], identifier + " progress")
            if progress is not None:
                if "%" in row[11].number_format:
                    progress *= 100
                if not 0 <= progress <= 100:
                    raise ValueError(f"{identifier}: progress outside 0..100")
                progress = round(progress, 2)
            records.append({"id": identifier, "name": by_id[identifier]["name"], "period": period,
                "date": report_date, "owner": txt(v[4]), "instance": instances[instance], "moment": moment,
                "execution": execution, "baseline": number(v[8], identifier + " baseline"),
                "target": number(v[9], identifier + " target"), "current": number(v[10], identifier + " current"),
                "progress": progress, "status": confidence[status], "comment": txt(v[13]), "evidence": txt(v[14]),
                "learning": txt(v[15]), "pending": txt(v[16]), "pending_treatment": txt(v[17]),
                "transfer_period": txt(v[18]), "strategy_learning": txt(v[19]),
                "source_sheet": tracking_sheet, "source_row": row[0].row, "source_moment": source_moment,
                "source_confidence": status, "source_instance": instance, "demo": any("[DEMO]" in txt(x) for x in v)})
        source = {"filename": input_path.name, "sha256": hashlib.sha256(input_path.read_bytes()).hexdigest()}
        generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
        strategy = {"metadata": {"publication": "public", "version": "v6", "kr_count": len(catalogue), "source": source, "generated_at": generated},
                    "objectives": sorted(objectives.values(), key=lambda o: o["number"]), "krs": catalogue}
        tracking = {"metadata": {"publication": "public", "version": "v6", "record_count": len(records),
                    "reported_kr_count": len({r["id"] for r in records}), "demo_records": sum(r["demo"] for r in records),
                    "periods": sorted({r["period"] for r in records}), "instances": sorted({r["instance"] for r in records}),
                    "source": source, "generated_at": generated, "moment_override": "Confirmado GE" if confirm_ge else None}, "records": records}
        return strategy, tracking, notes
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent / "data")
    parser.add_argument("--strategy-sheet", default="BD-Estrategia-Consolidada")
    parser.add_argument("--tracking-sheet", default="Seguimiento")
    parser.add_argument("--confirm-ge", action="store_true", help="Register all source reports as Confirmado GE when explicitly confirmed by the team.")
    args = parser.parse_args()
    strategy, tracking, notes = build(args.input, args.strategy_sheet, args.tracking_sheet, args.confirm_ge)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, data in [("strategy.json", strategy), ("tracking.json", tracking), ("objective-notes.json", notes)]:
        (args.output_dir / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"krs": len(strategy["krs"]), "reports": len(tracking["records"]),
                      "execution": dict(Counter(r["execution"] for r in tracking["records"])),
                      "demo_records": tracking["metadata"]["demo_records"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
