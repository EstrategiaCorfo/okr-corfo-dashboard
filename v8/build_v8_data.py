#!/usr/bin/env python3
"""Convert the current strategy workbook to the v8 dashboard JSON files."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re

import openpyxl
from openpyxl.utils import get_column_letter


def txt(value):
    return "" if value is None else str(value).strip()


def iso_date(value, label):
    if value is None or value == "":
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(txt(value), fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"{label}: invalid date {value!r}")


def periods(value):
    raw = txt(value)
    matches = list(re.finditer(r"(?:(20\d{2})\s*/\s*)?Q([1-4])", raw))
    if not matches or not matches[0].group(1):
        raise ValueError(f"Invalid target period: {raw!r}")
    year = int(matches[0].group(1))
    ordinals = []
    for match in matches:
        year = int(match.group(1) or year)
        ordinals.append(year * 4 + int(match.group(2)) - 1)
    if len(ordinals) == 2 and "-" in raw:
        if ordinals[1] < ordinals[0]:
            raise ValueError(f"Reversed target period: {raw!r}")
        ordinals = list(range(ordinals[0], ordinals[1] + 1))
    return list(dict.fromkeys(f"{p // 4}/Q{p % 4 + 1}" for p in ordinals))


def columns(sheet, header_row, fields):
    """Resolve by header text, keeping the rightmost canonical ID KR column."""
    found = {}
    for cell in sheet[header_row]:
        label = txt(cell.value).split("\n", 1)[0].strip()
        if label:
            found[label] = cell.column - 1
    missing = set(fields.values()) - found.keys()
    if missing:
        raise ValueError(f"{sheet.title}: missing columns {sorted(missing)}")
    return {key: found[label] for key, label in fields.items()}


STRATEGY_FIELDS = {
    "id": "ID KR", "dimension": "Dimensión", "objective": "Objetivo Estratégico",
    "vision": "Visión", "justification": "Justificación", "subobjective": "Subobjetivo estratégico",
    "name": "Nombre del KR", "target_period": "Trimestre meta KR", "target_year": "Año KR",
    "quarter": "Q", "target_date": "Fecha meta KR", "product": "Producto del KR (Referencial)",
    "area": "Área responsable", "owner": "Owner (Lider y responsable)",
    "substitutes": "Suplentes (apoyan ejecución)",
}
TRACKING_FIELDS = {
    "id": "ID KR", "objective": "Objetivo Estratégico", "subobjective": "Subobjetivo estratégico",
    "name": "Nombre del KR", "period": "Periodo de reporte", "date": "Fecha de reporte",
    "owner": "Owner", "instance": "Tipo de instancia", "moment": "Instancia del registro",
    "execution": "Estado de ejecución", "progress": "Métrica (%)",
    "confidence": "Nivel de confianza (semáforo)",
    "comment": "Comentario cualitativo (máximo 500 caracteres)", "evidence": "Evidencia / soporte",
    "learning": "Aprendizaje", "actions": "Acciones realizadas", "strategy_learning": "Aprendizaje",
    "pending": "¿Quedó pendiente de logro el KR?", "transfer_period": "Periodo de traspaso",
}


def build(input_path, strategy_sheet, tracking_sheet, confirm_ge):
    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    try:
        strategy_ws, tracking_ws = workbook[strategy_sheet], workbook[tracking_sheet]
        sc = columns(strategy_ws, 4, STRATEGY_FIELDS)
        tc = columns(tracking_ws, 2, TRACKING_FIELDS)
        # Both learning columns begin with "Aprendizaje"; their full headers differ.
        learning_columns = {txt(c.value): c.column - 1 for c in tracking_ws[2] if c.value}
        tc["learning"] = next((i for h, i in learning_columns.items() if h.startswith("Aprendizaje\n(")), None)
        tc["strategy_learning"] = learning_columns.get("Aprendizaje\nEstrategia")
        if tc["learning"] is None or tc["strategy_learning"] is None:
            raise ValueError("Seguimiento: both learning headers must be present")
        id_columns = [c.column - 1 for c in strategy_ws[4] if txt(c.value) == "ID KR"]

        objectives, notes, catalogue, ids = {}, {}, [], set()
        for row in strategy_ws.iter_rows(min_row=5):
            v = [c.value for c in row]
            get = lambda key: txt(v[sc[key]])
            identifier = get("id")
            if not identifier:
                continue
            if not re.fullmatch(r"KR-\d{3,}", identifier) or identifier in ids:
                raise ValueError(f"Invalid or duplicate KR ID: {identifier}")
            if any(txt(v[i]) and txt(v[i]) != identifier for i in id_columns):
                raise ValueError(f"{identifier}: duplicate ID columns disagree")
            ids.add(identifier)
            objective = get("objective")
            match = re.match(r"Objetivo\s+(\d+)\s*:", objective)
            if not match:
                raise ValueError(f"{identifier}: invalid objective {objective!r}")
            obj = objectives.setdefault(objective, {
                "name": objective, "number": int(match.group(1)), "dimension": get("dimension"),
                "vision": get("vision"), "subobjectives": [],
            })
            subobjective = get("subobjective")
            if subobjective not in obj["subobjectives"]:
                obj["subobjectives"].append(subobjective)
            if objective in notes and notes[objective] != get("justification"):
                raise ValueError(f"{identifier}: inconsistent objective justification")
            notes[objective] = get("justification")
            catalogue.append({
                **{key: get(key) for key in STRATEGY_FIELDS if key != "justification"},
                "periods": periods(get("target_period")),
                "target_date": iso_date(v[sc["target_date"]], identifier + " target date"),
                "verification": "", "comments": "", "source_sheet": strategy_sheet,
                "source_row": row[0].row,
            })

        by_id = {k["id"]: k for k in catalogue}
        confidence = {"Alto (On Track)": "On Track", "Medio (Off Track)": "Off Track",
                      "Bajo (At Risk)": "At Risk", "No iniciado": "No iniciado", "": ""}
        instances = {"Apertura de Q": "Apertura", "Apertura": "Apertura", "Revisión intermedia": "Revisión intermedia",
                     "Cierre de Q": "Cierre", "Cierre": "Cierre", "Pendiente": "Pendiente"}
        records, seen = [], set()
        for row in tracking_ws.iter_rows(min_row=3):
            v = [c.value for c in row]
            get = lambda key: txt(v[tc[key]])
            identifier = get("id")
            if not identifier:
                continue
            if identifier not in by_id:
                raise ValueError(f"Reporting KR is absent from catalogue: {identifier}")
            for key in ("name", "objective", "subobjective"):
                if get(key) and " ".join(get(key).split()) != " ".join(by_id[identifier][key].split()):
                    raise ValueError(f"{identifier}: reporting {key} does not match catalogue")
            period, instance, source_moment = get("period"), get("instance"), get("moment")
            execution, status = get("execution"), get("confidence")
            if not re.fullmatch(r"20\d{2}/Q[1-4]", period):
                raise ValueError(f"{identifier}: invalid reporting period {period!r}")
            if instance not in instances or execution not in {"No iniciado", "En proceso", "Completado"} or status not in confidence:
                raise ValueError(f"{identifier}: invalid instance, execution or confidence")
            if source_moment not in {"Propuesta Owner", "Confirmado GE"}:
                raise ValueError(f"{identifier}: invalid record validation {source_moment!r}")
            moment = "Confirmado GE" if confirm_ge else source_moment
            report_date = iso_date(v[tc["date"]], identifier + " report date")
            key = (identifier, period, instances[instance], moment, report_date)
            if key in seen:
                raise ValueError(f"Duplicate reporting record: {key}")
            seen.add(key)
            raw_progress = v[tc["progress"]]
            progress = None
            if raw_progress is not None and raw_progress != "":
                if isinstance(raw_progress, bool):
                    raise ValueError(f"{identifier}: progress must be numeric")
                progress = float(raw_progress)
                if "%" in row[tc["progress"]].number_format:
                    progress *= 100
                if not 0 <= progress <= 100:
                    raise ValueError(f"{identifier}: progress outside 0..100")
                progress = round(progress, 2)
            records.append({
                "id": identifier, "name": by_id[identifier]["name"], "period": period, "date": report_date,
                "owner": get("owner"), "instance": instances[instance], "moment": moment,
                "execution": execution, "baseline": None, "target": None, "current": None,
                "progress": progress, "status": confidence[status],
                **{field: get(field) for field in ("comment", "evidence", "learning", "actions", "strategy_learning", "pending", "transfer_period")},
                "pending_treatment": "", "source_sheet": tracking_sheet, "source_row": row[0].row,
                "source_moment": source_moment, "source_confidence": status, "source_instance": instance,
                "demo": any("[DEMO]" in txt(x) for x in v),
            })
        source = {"filename": input_path.name, "sha256": hashlib.sha256(input_path.read_bytes()).hexdigest()}
        generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
        metadata = {"publication": "public", "version": "v8", "source": source, "generated_at": generated}
        strategy = {"metadata": {**metadata, "kr_count": len(catalogue), "source_sheet": strategy_sheet,
                                "source_columns": {k: get_column_letter(i + 1) for k, i in sc.items()}},
                    "objectives": sorted(objectives.values(), key=lambda o: o["number"]), "krs": catalogue}
        tracking = {"metadata": {**metadata, "record_count": len(records),
                    "reported_kr_count": len({r["id"] for r in records}), "demo_records": sum(r["demo"] for r in records),
                    "periods": sorted({r["period"] for r in records}), "instances": sorted({r["instance"] for r in records}),
                    "moment_override": "Confirmado GE" if confirm_ge else None, "source_sheet": tracking_sheet,
                    "source_columns": {k: get_column_letter(i + 1) for k, i in tc.items()}}, "records": records}
        return strategy, tracking, notes
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent / "data")
    parser.add_argument("--strategy-sheet", default="BD-Estrategia-Consolidada")
    parser.add_argument("--tracking-sheet", default="Seguimiento")
    parser.add_argument("--confirm-ge", action="store_true", help="Apply the team's explicit GE confirmation to all reports.")
    args = parser.parse_args()
    strategy, tracking, notes = build(args.input, args.strategy_sheet, args.tracking_sheet, args.confirm_ge)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, data in (("strategy.json", strategy), ("tracking.json", tracking), ("objective-notes.json", notes)):
        (args.output_dir / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"krs": len(strategy["krs"]), "reports": len(tracking["records"]),
                      "execution": dict(Counter(r["execution"] for r in tracking["records"])),
                      "confidence": dict(Counter(r["status"] for r in tracking["records"])),
                      "actions": sum(bool(r["actions"]) for r in tracking["records"]),
                      "demo_records": tracking["metadata"]["demo_records"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
