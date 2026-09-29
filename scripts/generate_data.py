#!/usr/bin/env python3
"""Valida la planilla maestra y genera JSON locales o para el sitio público."""

import argparse
import json
import re
import sys
import warnings
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


STRATEGY_COLUMNS = {
    'dimension': 'Dimensión',
    'objective': 'Objetivo Estratégico',
    'vision': 'Visión',
    'subobjective': 'Subobjetivo estratégico',
    'id': 'ID KR',
    'name': 'Nombre del KR',
    'target_period': 'Trimestre meta KR',
    'target_year': 'Año KR',
    'quarter': 'Q',
    'target_date': 'Fecha meta KR',
    'product': 'Producto asociado estimado',
    'verification': 'Medio de verificación OKR estimado',
    'area': 'Área quien lidera',
    'owner': 'Owner (Lider y responsable)',
    'substitutes': 'Suplentes (apoyan ejecución)',
    'comments': 'Comentarios',
    'kr_type': 'Tipo de KR',
    'cdc': 'Asociado a CDC',
}
TRACKING_COLUMNS = {
    'id': 'ID KR',
    'name': 'Nombre del KR',
    'period': 'Periodo de reporte',
    'date': 'Fecha de reporte',
    'instance': 'Instancia',
    'moment': 'Momento del registro',
    'execution': 'Estado de ejecución',
    'baseline': 'Línea base',
    'target': 'Meta',
    'current': 'Valor actual',
    'progress': 'Progreso %',
    'status': 'Estatus (semáforo)',
    'comment': 'Comentario cualitativo',
    'evidence': 'Evidencia / soporte',
    'learning': 'Aprendizaje',
    'pending': '¿Quedó pendiente?',
    'pending_treatment': 'Tratamiento del pendiente',
    'transfer_period': 'Periodo de traspaso',
}
COLUMN_ALIASES = {
    'Estrategia': {
        'target_period': ('Trimestre meta KR (inicial)',),
    },
    'Seguimiento': {
        'execution': ('Estado del hito',),
        'moment': ('Instancia del registro',),
        'progress': ('Progreso numérico (%)', 'Progreso numérico', 'Métrica (%)'),
        'status': ('Nivel de confianza', 'Nivel de confianza (semáforo)'),
        'comment': ('Comentario cualitativo (máximo 500 caracteres)',),
    },
}
REQUIRED_STRATEGY = {'dimension', 'objective', 'vision', 'subobjective', 'id', 'name', 'target_period', 'target_year', 'quarter', 'target_date', 'area', 'owner'}
REQUIRED_TRACKING = {'id', 'period', 'date', 'instance', 'moment', 'execution', 'baseline', 'target', 'current', 'progress', 'status', 'comment', 'evidence', 'learning'}
EXECUTIONS = {'No iniciado', 'En proceso', 'Completado'}
STATUSES = {
    'En Cumplimiento / On Track': 'On Track',
    'Alerta de cumplimiento / Off Track': 'Off Track',
    'Riesgo de cumplimiento / At Risk': 'At Risk',
    'On Track': 'On Track', 'Off Track': 'Off Track', 'At Risk': 'At Risk',
    'Alto (On Track)': 'On Track', 'Medio (Off Track)': 'Off Track',
    'Bajo (At Risk)': 'At Risk',
    'Alto': 'On Track', 'Medio': 'Off Track', 'Bajo': 'At Risk',
    'Alto / On Track': 'On Track', 'Medio / Off Track': 'Off Track',
    'Bajo / At Risk': 'At Risk',
}
MOMENTS = {'Confirmado GE', 'Propuesta Owner'}
PERIOD = re.compile(r'^(20\d{2})/Q([1-4])$')
RANGE_SAME_YEAR = re.compile(r'^(20\d{2})/Q([1-4])-Q([1-4])$')
RANGE_YEARS = re.compile(r'^(20\d{2})/Q([1-4])\s*-\s*(20\d{2})/Q([1-4])$')
def txt(value):
    return '' if value is None else str(value).strip()


def date_iso(value, epoch):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value, epoch)
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, str):
        for fmt in ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d'):
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    raise ValueError('fecha inválida (usar fecha Excel o DD-MM-AAAA)')


def numeric(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        raise ValueError('valor numérico inválido')
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(txt(value).replace(' ', '').replace(',', '.'))
    except ValueError as exc:
        raise ValueError('valor numérico inválido') from exc


def periods(value):
    value = txt(value)
    if match := PERIOD.fullmatch(value):
        return [value]
    if match := RANGE_SAME_YEAR.fullmatch(value):
        year, first, last = map(int, match.groups())
        if last >= first:
            return [f'{year}/Q{q}' for q in range(first, last + 1)]
    if match := RANGE_YEARS.fullmatch(value):
        y1, q1, y2, q2 = map(int, match.groups())
        first, last = y1 * 4 + q1 - 1, y2 * 4 + q2 - 1
        if first <= last and last - first < 21:
            return [f'{x // 4}/Q{x % 4 + 1}' for x in range(first, last + 1)]
    raise ValueError('periodo inválido (usar AAAA/Qn, AAAA/Qn-Qn o AAAA/Qn - AAAA/Qn)')


def headers(sheet, row, required, definitions, errors):
    cells = [txt(c.value) for c in sheet[row]]
    aliases = COLUMN_ALIASES.get(sheet.title, {})
    found = {}
    for key, label in definitions.items():
        matches = [cells.index(candidate) for candidate in (label, *aliases.get(key, ())) if candidate in cells]
        if len(matches) > 1:
            errors.append(f'{sheet.title}, encabezado fila {row}: hay dos columnas para «{label}»')
        if matches:
            found[key] = matches[0]
    for key in sorted(required - found.keys()):
        accepted = ', '.join(f'«{x}»' for x in (definitions[key], *aliases.get(key, ())))
        errors.append(f'{sheet.title}, encabezado fila {row}: falta columna {accepted}')
    if len(cells) != len(set(c for c in cells if c)):
        errors.append(f'{sheet.title}, encabezado fila {row}: hay columnas duplicadas')
    return found


def value(row, mapping, key):
    return row[mapping[key]].value if key in mapping else None


def reported_progress(cell):
    """Lee un porcentaje escrito directamente, incluso si Excel lo almacena como fracción."""
    raw = cell.value
    if raw is None or raw == '' or (isinstance(raw, str) and raw.startswith('=')):
        return None
    if isinstance(raw, str) and raw.strip().endswith('%'):
        return numeric(raw.strip()[:-1])
    result = numeric(raw)
    return result * 100 if '%' in cell.number_format else result


def build(path):
    warnings.filterwarnings('ignore', message='Unknown extension is not supported and will be removed')
    warnings.filterwarnings('ignore', message='Conditional Formatting extension is not supported and will be removed')
    wb = load_workbook(path, read_only=True, data_only=False)
    errors = []
    if set(wb.sheetnames) != {'Estrategia', 'Seguimiento'}:
        raise ValueError('La planilla debe contener exactamente las hojas Estrategia y Seguimiento.')
    ws, wt = wb['Estrategia'], wb['Seguimiento']
    hs = headers(ws, 4, REQUIRED_STRATEGY, STRATEGY_COLUMNS, errors)
    ht = headers(wt, 1, REQUIRED_TRACKING, TRACKING_COLUMNS, errors)
    if errors:
        raise ValueError('\n'.join(errors))
    krs, by_id, by_name, objectives = [], {}, {}, {}
    for number, row in enumerate(ws.iter_rows(min_row=5), 5):
        ident = txt(value(row, hs, 'id'))
        if not ident:
            continue
        location = f'Estrategia, fila {number} ({ident})'
        if not re.fullmatch(r'KR-\d{3,}', ident):
            errors.append(f'{location}: ID KR inválido; usar KR-001, KR-002, etc.')
        if ident in by_id:
            errors.append(f'{location}: ID KR duplicado (ya existe en fila {by_id[ident]})')
        by_id[ident] = number
        try:
            active = periods(value(row, hs, 'target_period'))
        except ValueError as exc:
            errors.append(f'{location}, Trimestre meta KR: {exc}')
            active = []
        try:
            due = date_iso(value(row, hs, 'target_date'), wb.epoch)
        except ValueError as exc:
            errors.append(f'{location}, Fecha meta KR: {exc}')
            due = None
        for key in ('dimension', 'objective', 'vision', 'subobjective', 'name'):
            if not txt(value(row, hs, key)):
                errors.append(f'{location}: falta «{STRATEGY_COLUMNS[key]}»')
        objective = txt(value(row, hs, 'objective'))
        match = re.match(r'^Objetivo\s+(\d+)\s*:\s*(.+)$', objective, re.I)
        if not match:
            errors.append(f'{location}: Objetivo Estratégico debe comenzar con «Objetivo N:»')
        year = value(row, hs, 'target_year')
        if active and (str(year) != active[0][:4] or txt(value(row, hs, 'quarter'))[:2] != active[0][-2:]):
            errors.append(f'{location}: Año KR o Q no coincide con Trimestre meta KR')
        obj_key = objective
        info = {
            'id': ident,
            'name': txt(value(row, hs, 'name')),
            'objective': obj_key,
            'subobjective': txt(value(row, hs, 'subobjective')),
            'dimension': txt(value(row, hs, 'dimension')),
            'vision': txt(value(row, hs, 'vision')),
            'target_period': txt(value(row, hs, 'target_period')),
            'target_year': txt(year),
            'quarter': txt(value(row, hs, 'quarter')),
            'periods': active,
            'target_date': due,
            'product': txt(value(row, hs, 'product')),
            'verification': txt(value(row, hs, 'verification')),
            'area': txt(value(row, hs, 'area')),
            'owner': txt(value(row, hs, 'owner')),
            'substitutes': txt(value(row, hs, 'substitutes')),
            'comments': txt(value(row, hs, 'comments')),
            'kr_type': txt(value(row, hs, 'kr_type')) or None,
            'cdc': None,
        }
        if info['kr_type'] and info['kr_type'] not in {'Resultado final', 'Paso intermedio'}:
            errors.append(f'{location}, Tipo de KR: usar Resultado final o Paso intermedio')
        cdc_value = txt(value(row, hs, 'cdc')).lower()
        if cdc_value in {'sí', 'si', '1', 'true'}:
            info['cdc'] = True
        elif cdc_value in {'no', '0', 'false'}:
            info['cdc'] = False
        elif cdc_value:
            errors.append(f'{location}, Asociado a CDC: usar Sí o No')
        krs.append(info)
        by_name[ident] = info['name']
        if obj_key and obj_key not in objectives:
            objectives[obj_key] = {
                'name': obj_key,
                'number': int(match[1]) if match else None,
                'dimension': info['dimension'],
                'vision': info['vision'],
                'subobjectives': [],
            }
        if obj_key:
            obj = objectives[obj_key]
            if info['subobjective'] not in obj['subobjectives']:
                obj['subobjectives'].append(info['subobjective'])
            if (info['dimension'], info['vision']) != (obj['dimension'], obj['vision']):
                errors.append(f'{location}: dimensión o visión contradice otra fila del mismo objetivo')

    records = []
    for number, row in enumerate(wt.iter_rows(min_row=2), 2):
        ident = txt(value(row, ht, 'id'))
        if not ident:
            continue
        location = f'Seguimiento, fila {number} ({ident})'
        if ident not in by_id:
            errors.append(f'{location}: ID KR no existe en Estrategia')
        raw_period = txt(value(row, ht, 'period'))
        if not PERIOD.fullmatch(raw_period):
            errors.append(f'{location}, Periodo de reporte: usar AAAA/Qn')
        try:
            report_date = date_iso(value(row, ht, 'date'), wb.epoch)
        except ValueError as exc:
            errors.append(f'{location}, Fecha de reporte: {exc}')
            report_date = None
        execution = txt(value(row, ht, 'execution'))
        status = txt(value(row, ht, 'status'))
        moment = txt(value(row, ht, 'moment'))
        if execution and execution not in EXECUTIONS:
            errors.append(f'{location}, Estado de ejecución: «{execution}» no es válido')
        if status and status not in STATUSES:
            errors.append(f'{location}, Nivel de confianza: «{status}» no es válido')
        if moment not in MOMENTS:
            errors.append(f'{location}, Momento del registro: usar Confirmado GE o Propuesta Owner')
        vals = []
        for key in ('baseline', 'target', 'current'):
            try:
                vals.append(numeric(value(row, ht, key)))
            except ValueError as exc:
                errors.append(f'{location}, {TRACKING_COLUMNS[key]}: {exc}')
                vals.append(None)
        if 0 < sum(v is not None for v in vals) < 3:
            errors.append(f'{location}: Línea base, Meta y Valor actual deben informarse juntos')
        progress = None
        if all(v is not None for v in vals):
            if vals[1] == vals[0]:
                errors.append(f'{location}: Meta no puede ser igual a Línea base')
            else:
                progress = (vals[2] - vals[0]) / (vals[1] - vals[0]) * 100
        if 'progress' in ht:
            try:
                given = reported_progress(row[ht['progress']])
                if given is not None:
                    # Métrica (%) es el avance informado por el owner, incluso
                    # para KR cualitativos con valores de referencia 0 y 1.
                    if txt(wt.cell(1, ht['progress'] + 1).value) == 'Métrica (%)':
                        progress = given
                    elif progress is not None and abs(given - progress) > .11:
                        errors.append(f'{location}, Progreso: difiere del cálculo de Línea base, Meta y Valor actual')
                    elif progress is None:
                        progress = given
            except ValueError as exc:
                errors.append(f'{location}, Progreso: {exc}')
        if 'name' in ht and ident in by_id:
            given_name = value(row, ht, 'name')
            expected = by_name[ident]
            if given_name is not None and not (isinstance(given_name, str) and given_name.startswith('=')) and txt(given_name) != expected:
                errors.append(f'{location}: Nombre del KR difiere de Estrategia')
        records.append({
            'id': ident, 'period': raw_period, 'date': report_date,
            'instance': txt(value(row, ht, 'instance')),
            'moment': moment, 'execution': execution or None,
            'baseline': vals[0], 'target': vals[1], 'current': vals[2],
            'progress': round(progress, 2) if progress is not None else None,
            'status': STATUSES.get(status),
            'comment': txt(value(row, ht, 'comment')),
            'evidence': txt(value(row, ht, 'evidence')),
            'learning': txt(value(row, ht, 'learning')),
            'pending': txt(value(row, ht, 'pending')),
            'pending_treatment': txt(value(row, ht, 'pending_treatment')),
            'transfer_period': txt(value(row, ht, 'transfer_period')),
            'source_row': number,
        })
    if not krs:
        errors.append('Estrategia: no se encontraron KR')
    if errors:
        raise ValueError('\n'.join(errors))
    demo_count = sum('[DEMO]' in ' '.join(str(v) for v in record.values()) for record in records)
    strategy = {'metadata': {'source': path.name, 'kr_count': len(krs)},
                'objectives': sorted(objectives.values(), key=lambda item: item['number'] or 999),
                'krs': krs}
    tracking = {'metadata': {'source': path.name, 'record_count': len(records), 'demo_records': demo_count},
                'records': records}
    return strategy, tracking


def normalize_strategy(strategy, catalog_path):
    """Aplica el catálogo oficial sin alterar los ID de KR ni la planilla original."""
    catalog = json.loads(catalog_path.read_text(encoding='utf-8'))['objectives']
    if len(catalog) != 8 or {item['number'] for item in catalog} != set(range(1, 9)):
        raise ValueError('El catálogo debe contener los ocho objetivos oficiales.')
    source_map = {item['source_name']: item for item in catalog}
    current_map = {item['name']: item for item in catalog}
    if len(source_map) != 8 or len(current_map) != 8:
        raise ValueError('El catálogo contiene nombres de objetivo duplicados.')
    official_dimensions = {'Impulso Corfo', 'Rol de Corfo', 'Habilitantes', 'Impacto'}
    if {item['dimension'] for item in catalog} != official_dimensions:
        raise ValueError('El catálogo contiene una dimensión fuera de las cuatro oficiales.')
    for kr in strategy['krs']:
        entry = source_map.get(kr['objective']) or current_map.get(kr['objective'])
        if not entry:
            raise ValueError(f"{kr['id']}: objetivo ausente del catálogo oficial: {kr['objective']}")
        if kr['objective'] in source_map and kr['dimension'] != entry['source_dimension']:
            raise ValueError(f"{kr['id']}: dimensión de origen no coincide con el catálogo.")
        if kr['objective'] in current_map and kr['dimension'] != entry['dimension']:
            raise ValueError(f"{kr['id']}: dimensión oficial inválida.")
        match = re.match(r'^SO(\d+)\.(\d+)(?:\b|:)', kr['subobjective'])
        expected = entry['old_number'] if kr['objective'] in source_map else entry['number']
        if not match or int(match[1]) != expected or not 1 <= int(match[2]) <= len(entry['subobjectives']):
            raise ValueError(f"{kr['id']}: subobjetivo ausente del catálogo oficial: {kr['subobjective']}")
        kr.update({
            'objective': entry['name'], 'subobjective': entry['subobjectives'][int(match[2]) - 1],
            'dimension': entry['dimension'], 'vision': entry['vision'],
        })
    strategy['objectives'] = [{
        'name': item['name'], 'number': item['number'], 'dimension': item['dimension'],
        'vision': item['vision'], 'subobjectives': item['subobjectives'],
    } for item in catalog]
    return {item['name']: item['justification'] for item in catalog}, source_map


def public_view(strategy, tracking):
    """Publica los campos de la planilla y marca cada reporte de demostración."""
    records = []
    for record in tracking['records']:
        published_record = record.copy()
        published_record['demo'] = any('[DEMO]' in str(value) for value in record.values())
        records.append(published_record)
    public_strategy = {
        'metadata': {'publication': 'public', 'kr_count': len(strategy['krs'])},
        'objectives': strategy['objectives'],
        'krs': strategy['krs'],
    }
    public_tracking = {
        'metadata': {
            'publication': 'public', 'record_count': len(records),
            'demo_records': sum(record['demo'] for record in records),
        },
        'records': records,
    }
    return public_strategy, public_tracking


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path, help='Ruta del Excel con hojas Estrategia y Seguimiento')
    parser.add_argument('--output', type=Path, help='Carpeta de salida (por defecto data/ o data/public/ con --public)')
    parser.add_argument('--validate-only', action='store_true', help='No escribir JSON')
    parser.add_argument('--require-real', action='store_true', help='Rechazar filas marcadas [DEMO]')
    parser.add_argument('--public', action='store_true', help='Generar JSON públicos con todos los campos de la planilla')
    parser.add_argument('--notes', type=Path, help='Justificaciones opcionales; por defecto se usa el catálogo oficial')
    parser.add_argument('--catalog', type=Path, default=Path(__file__).resolve().parent.parent / 'data/strategy-catalog.json', help='Catálogo oficial de objetivos y justificaciones')
    args = parser.parse_args()
    try:
        strategy, tracking = build(args.input)
        notes, source_map = normalize_strategy(strategy, args.catalog)
        if args.require_real and tracking['metadata']['demo_records']:
            raise ValueError(f"Seguimiento: {tracking['metadata']['demo_records']} registros contienen [DEMO]; se requieren datos reales.")
        if args.notes:
            overrides = json.loads(args.notes.read_text(encoding='utf-8'))
            if not isinstance(overrides, dict):
                raise ValueError('Las justificaciones deben ser un objeto JSON.')
            remapped = {entry['name']: overrides.get(entry['name'], overrides.get(source_name)) for source_name, entry in source_map.items()}
            if set(remapped) != set(notes) or not all(isinstance(text, str) and text.strip() for text in remapped.values()):
                raise ValueError('Las justificaciones deben ser textos no vacíos para todos los objetivos de la planilla.')
            notes = remapped
    except (OSError, ValueError) as exc:
        print(f'ERROR de validación:\n{exc}', file=sys.stderr)
        return 1
    print(f"Validación correcta: {len(strategy['krs'])} KR, {len(tracking['records'])} reportes, "
          f"{len(set(r['id'] for r in tracking['records']))} KR reportados.")
    print(f"Sin Owner informado: {sum(not k['owner'] for k in strategy['krs'])} KR.")
    if tracking['metadata']['demo_records']:
        print(f"ADVERTENCIA: {tracking['metadata']['demo_records']} reportes están marcados [DEMO]. No presentarlos como avances reales.")
    if not args.validate_only:
        if args.public:
            strategy, tracking = public_view(strategy, tracking)
            print('Vista pública: incluidos responsables y campos cualitativos de la planilla; reportes DEMO identificados.')
        args.output = args.output or Path('data/public' if args.public else 'data')
        args.output.mkdir(parents=True, exist_ok=True)
        for name, content in (('strategy.json', strategy), ('tracking.json', tracking)):
            target = args.output / name
            tmp = target.with_suffix('.json.tmp')
            tmp.write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
            tmp.replace(target)
            print(f'Generado: {target}')
        target = args.output / 'objective-notes.json'
        tmp = target.with_suffix('.json.tmp')
        tmp.write_text(json.dumps(notes, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        tmp.replace(target)
        print(f'Generado: {target}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
