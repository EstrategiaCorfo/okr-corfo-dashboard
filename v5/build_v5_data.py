#!/usr/bin/env python3
"""Exporta Estrategia y Seguimiento_KR de la planilla v5, sin publicar los ejemplos."""

import argparse
import json
import re
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'scripts'))
from generate_data import date_iso, numeric, periods, txt, STATUSES, EXECUTIONS, MOMENTS  # noqa: E402

STRATEGY = {
    'dimension': 'Dimensión', 'objective': 'Objetivo Estratégico', 'vision': 'Visión',
    'justification': 'Justificación', 'subobjective': 'Subobjetivo estratégico',
    'id': 'ID KR', 'name': 'Nombre del KR', 'target_period': 'Trimestre meta KR',
    'target_year': 'Año KR', 'quarter': 'Q', 'target_date': 'Fecha meta KR',
    'product': 'Producto asociado al KR (Referencial)', 'area': 'Área responsable',
    'owner': 'Owner (Lider y responsable)', 'substitutes': 'Suplentes (apoyan ejecución)',
}
TRACKING = {
    'id': 'ID KR', 'name': 'Nombre del KR', 'period': 'Periodo de reporte',
    'date': 'Fecha de reporte', 'owner': 'Owner', 'instance': 'Instancia',
    'moment': 'Instancia del registro', 'execution': 'Estado de ejecución',
    'baseline': 'Línea base', 'target': 'Meta', 'current': 'Valor actual',
    'progress': 'Métrica (%)', 'status': 'Nivel de confianza (semáforo)',
    'comment': 'Comentario cualitativo (máximo 500 caracteres)',
    'evidence': 'Evidencia / soporte', 'learning': 'Aprendizaje',
    'pending': '¿Quedó pendiente de logro el KR?', 'pending_treatment': 'Tratamiento del pendiente',
    'transfer_period': 'Periodo de traspaso',
}
INSTANCES = {'Apertura': 'Apertura', 'Apertura de Q': 'Apertura',
             'Revisión intermedia': 'Revisión intermedia', 'Cierre': 'Cierre', 'Cierre de Q': 'Cierre'}


def columns(sheet, row, definitions):
    labels = [txt(cell.value).split('\n', 1)[0].strip() for cell in sheet[row]]
    if len([v for v in labels if v]) != len(set(v for v in labels if v)):
        raise ValueError(f'{sheet.title}: encabezados duplicados')
    missing = set(definitions.values()) - set(labels)
    if missing:
        raise ValueError(f'{sheet.title}: faltan columnas {sorted(missing)}')
    return {key: labels.index(label) for key, label in definitions.items()}


def build(path):
    warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
    wb = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    if not {'Estrategia', 'Seguimiento_KR'} <= set(wb.sheetnames):
        raise ValueError('Se requieren las hojas Estrategia y Seguimiento_KR')
    ws, wt = wb['Estrategia'], wb['Seguimiento_KR']
    hs, ht = columns(ws, 4, STRATEGY), columns(wt, 2, TRACKING)
    krs, by_id, objectives = [], {}, {}
    for number, row in enumerate(ws.iter_rows(min_row=5), 5):
        kr = {key: txt(row[col].value) for key, col in hs.items()}
        if not kr['id']:
            continue
        if not re.fullmatch(r'KR-\d{3,}', kr['id']) or kr['id'] in by_id:
            raise ValueError(f'Estrategia, fila {number}: ID inválido o duplicado')
        match = re.fullmatch(r'Objetivo\s+(\d+)\s*:\s*(.+)', kr['objective'])
        if not match or not kr['name'] or not kr['vision'] or not kr['subobjective']:
            raise ValueError(f'Estrategia, fila {number}: definición incompleta')
        active = periods(kr['target_period'])
        if kr['target_year'] != active[0][:4] or kr['quarter'][:2] != active[0][-2:]:
            raise ValueError(f'Estrategia, fila {number}: año o trimestre inconsistente')
        kr.update(periods=active, target_date=date_iso(row[hs['target_date']].value, wb.epoch),
                  kr_type=None, cdc=None, source_row=number)
        krs.append(kr)
        by_id[kr['id']] = kr
        obj = objectives.setdefault(kr['objective'], {
            'name': kr['objective'], 'number': int(match[1]), 'dimension': kr['dimension'],
            'vision': kr['vision'], 'justification': kr['justification'], 'subobjectives': [],
        })
        if any(obj[key] != kr[key] for key in ('dimension', 'vision', 'justification')):
            raise ValueError(f'Estrategia, fila {number}: definición de objetivo inconsistente')
        if kr['subobjective'] not in obj['subobjectives']:
            obj['subobjectives'].append(kr['subobjective'])

    records = []
    for number, row in enumerate(wt.iter_rows(min_row=3), 3):
        ident = txt(row[ht['id']].value)
        if not ident:
            continue
        if ident not in by_id:
            raise ValueError(f'Seguimiento_KR, fila {number}: KR no definido en Estrategia')
        record = {}
        for key, col in ht.items():
            cell = row[col]
            raw = cell.value
            if cell.data_type == 'f':
                raw = cached['Seguimiento_KR'].cell(number, col + 1).value
                if raw is None and key != 'name':
                    raise ValueError(f'Seguimiento_KR, {cell.coordinate}: fórmula sin resultado guardado')
            record[key] = raw
        given_name = txt(record.pop('name'))
        if given_name and given_name != by_id[ident]['name']:
            raise ValueError(f'Seguimiento_KR, fila {number}: nombre de KR inconsistente')
        period = txt(record['period'])
        if not re.fullmatch(r'20\d{2}/Q[1-4]', period):
            raise ValueError(f'Seguimiento_KR, fila {number}: periodo inválido')
        if txt(record['instance']) not in INSTANCES:
            raise ValueError(f'Seguimiento_KR, fila {number}: instancia inválida')
        if txt(record['moment']) not in MOMENTS:
            raise ValueError(f'Seguimiento_KR, fila {number}: instancia del registro inválida')
        if txt(record['execution']) not in EXECUTIONS | {''}:
            raise ValueError(f'Seguimiento_KR, fila {number}: estado de ejecución inválido')
        status = txt(record['status'])
        if status not in set(STATUSES) | {'No iniciado', ''}:
            raise ValueError(f'Seguimiento_KR, fila {number}: nivel de confianza inválido')
        for key in ('baseline', 'target', 'current'):
            record[key] = numeric(record[key])
        metric = record['progress']
        if isinstance(metric, str) and metric.strip().endswith('%'):
            progress = numeric(metric.strip()[:-1])
        else:
            progress = numeric(metric)
            if progress is not None and '%' in row[ht['progress']].number_format:
                progress *= 100
        values = [record[key] for key in ('baseline', 'target', 'current')]
        calculated = ((values[2] - values[0]) / (values[1] - values[0]) * 100
                      if all(v is not None for v in values) and values[1] != values[0] else None)
        basis = 'reported' if progress is not None else 'calculated' if calculated is not None else None
        if progress is None:
            progress = calculated
        if progress is not None and not 0 <= progress <= 100:
            raise ValueError(f'Seguimiento_KR, fila {number}: métrica fuera de 0 a 100%')
        for key in ('owner', 'moment', 'comment', 'evidence', 'learning', 'pending', 'pending_treatment', 'transfer_period'):
            record[key] = txt(record[key])
        record.update(id=ident, period=period, date=date_iso(record['date'], wb.epoch),
                      instance=INSTANCES[txt(record['instance'])], execution=txt(record['execution']) or None,
                      status='No iniciado' if status == 'No iniciado' else STATUSES.get(status),
                      progress=round(progress, 2) if progress is not None else None,
                      progress_basis=basis, source_status=status, source_instance=txt(row[ht['instance']].value),
                      source_sheet='Seguimiento_KR', source_row=number)
        record['demo'] = any('[DEMO]' in str(value) for value in record.values())
        records.append(record)
    if not krs:
        raise ValueError('Estrategia no contiene KR')
    reported = len({r['id'] for r in records})
    strategy = {'metadata': {'publication': 'public', 'version': 'v5', 'visibility': 'full-catalog',
                            'kr_count': len(krs), 'catalog_kr_count': len(krs), 'reported_kr_count': reported,
                            'source_sheet': 'Estrategia'},
                'objectives': sorted(objectives.values(), key=lambda o: o['number']), 'krs': krs}
    tracking = {'metadata': {'publication': 'public', 'version': 'v5', 'source_sheet': 'Seguimiento_KR',
                            'excluded_sheets': [name for name in wb.sheetnames if name not in {'Estrategia', 'Seguimiento_KR'}],
                            'record_count': len(records), 'reported_kr_count': reported,
                            'demo_records': sum(r['demo'] for r in records)}, 'records': records}
    notes = {obj['name']: obj['justification'] for obj in strategy['objectives']}
    return strategy, tracking, notes


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path)
    args = parser.parse_args()
    strategy, tracking, notes = build(args.input)
    folder = Path(__file__).resolve().parent / 'data'
    folder.mkdir(exist_ok=True)
    for name, payload in [('strategy.json', strategy), ('tracking.json', tracking), ('objective-notes.json', notes)]:
        (folder / name).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f"v5: {len(strategy['krs'])} KR, {len(tracking['records'])} registros de Seguimiento_KR, "
          f"{tracking['metadata']['demo_records']} registros DEMO; ejemplos excluidos.")


if __name__ == '__main__':
    main()
